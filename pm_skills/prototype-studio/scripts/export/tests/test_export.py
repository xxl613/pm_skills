import base64,importlib.util,json,shutil,sys,tempfile,unittest,zipfile
from pathlib import Path
EXPORT=Path(__file__).resolve().parents[1]
def load(name):
 spec=importlib.util.spec_from_file_location(name,EXPORT/(name+'.py'));m=importlib.util.module_from_spec(spec);spec.loader.exec_module(m);return m
bundle=load('package_prototype');word=load('export_prd_docx')
class ExportTests(unittest.TestCase):
 def setUp(self):
  self.temp=tempfile.TemporaryDirectory();self.root=Path(self.temp.name);self.src=self.root/'source';self.src.mkdir()
 def tearDown(self):self.temp.cleanup()
 def write(self,name,text):
  p=self.src/name;p.parent.mkdir(parents=True,exist_ok=True);p.write_text(text,encoding='utf-8');return p
 def test_static_multifile_prd_navigation_and_zip(self):
  self.write('index.html','<link href="a.css" rel="stylesheet"><link rel="stylesheet" href="b.css"><script src="a.js"></script><script src="b.js"></script><a href="prd/学习.md">PRD</a>')
  self.write('a.css','@import "c.css";body{background:url("img/x.svg")}');self.write('b.css','p{color:red}');self.write('c.css','body{margin:0}');self.write('img/x.svg','<svg xmlns="http://www.w3.org/2000/svg"/>')
  self.write('a.js','// fetch("not-real.json")\nwindow.pages=[{path:"second.html"}];');self.write('b.js','window.ready=true;');self.write('second.html','<h1>学习</h1>');self.write('prd/学习.md','# 页面PRD\n\n全文。')
  result=bundle.package(self.src,self.root/'bundle',zip_path=self.root/'share.zip')
  self.assertEqual(result['brokenReferences'],[]);self.assertFalse(result['browserVerified']);self.assertTrue((self.root/'bundle/prd/学习.md').is_file())
  self.assertEqual((self.src/'index.html').read_bytes(),(self.root/'bundle/index.html').read_bytes())
  with zipfile.ZipFile(self.root/'share.zip') as z:self.assertIsNone(z.testzip());self.assertIn('bundle/b.js',z.namelist())
 def test_js_missing_page_detected(self):
  self.write('index.html','<script src="nav.js"></script>');self.write('nav.js','const items=[{path:"missing.html"}];')
  with self.assertRaisesRegex(ValueError,'missing.html'):bundle.package(self.src,self.root/'out')
  self.assertFalse((self.root/'out').exists())
 def test_vite_split_modules_server_only(self):
  self.write('index.html','<script type="module" src="assets/app.js"></script>');self.write('assets/app.js','import("./chunk.js").then(console.log);');self.write('assets/chunk.js','export const x=1;')
  with self.assertRaisesRegex(ValueError,'--mode server'):bundle.package(self.src,self.root/'direct')
  r=bundle.package(self.src,self.root/'server',mode='server');self.assertTrue(r['runtimeRisks']);self.assertTrue((self.root/'server/serve.py').is_file());self.assertTrue((self.root/'server/assets/chunk.js').is_file())
 def test_inline_fetch_requires_server(self):
  self.write('index.html','<script>fetch("data.json")</script>');self.write('data.json','{}')
  with self.assertRaisesRegex(ValueError,'source-origin'):bundle.package(self.src,self.root/'out')
 def test_comments_and_string_text_not_fetch_calls(self):
  self.write('index.html','<script>/* fetch("x.json") */ const help="fetch is not used"; window.ok=true;</script>')
  self.assertEqual(bundle.package(self.src,self.root/'out')['runtimeRisks'],[])
 def test_url_constructor_not_css_resource(self):
  self.write('index.html','<script>const current=new URL(location.href); window.scene=current.searchParams.get("scene");</script>')
  r=bundle.package(self.src,self.root/'out');self.assertEqual(r['brokenReferences'],[]);self.assertEqual(r['runtimeRisks'],[])
 def test_existing_output_untouched(self):
  self.write('index.html','test');out=self.root/'out';out.mkdir();(out/'keep').write_text('user')
  with self.assertRaisesRegex(ValueError,'exists'):bundle.package(self.src,out)
  self.assertEqual((out/'keep').read_text(),'user')
 def test_external_assets_not_offline(self):
  self.write('index.html','<img src="https://example.com/a.png">')
  with self.assertRaisesRegex(ValueError,'direct-open'):bundle.package(self.src,self.root/'direct')
  self.assertEqual(len(bundle.package(self.src,self.root/'server',mode='server')['remoteResources']),1)
 def test_symlink_rejected(self):
  self.write('index.html','test');(self.src/'linked.md').symlink_to(self.root/'somewhere')
  with self.assertRaisesRegex(ValueError,'symlink'):bundle.package(self.src,self.root/'out')
 def test_incomplete_review_contract_rejected(self):
  self.write('index.html','test');self.write('pages.json','{}')
  with self.assertRaisesRegex(ValueError,'Missing source contract'):bundle.package(self.src,self.root/'out')
 @unittest.skipUnless(shutil.which('pandoc'),'Pandoc unavailable')
 def test_word_real_content_and_style(self):
  image=self.src/'dot.png';image.write_bytes(base64.b64decode('iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+jFioAAAAASUVORK5CYII='))
  source=self.write('页面 PRD.md','# 页面标题\n\n中文正文，**保留强调**。\n\n## 功能\n\n- 新建\n- 修改\n\n| 字段 | 说明 |\n|---|---|\n| ID | 内容较长的业务说明，用于检查列宽。 |\n\n![image-202601020304](dot.png)\n')
  r=word.export(source,self.root/'prd.docx');self.assertTrue(r['structureVerified']);self.assertFalse(r['visualVerified']);self.assertEqual(r['docxContentCounts'],{'Header':2,'Table':1,'Image':1})
  from docx import Document
  d=Document(self.root/'prd.docx');self.assertEqual(d.paragraphs[0].text,'页面标题');self.assertTrue(any('中文正文' in p.text for p in d.paragraphs));self.assertNotEqual(d.tables[0].columns[0].width,d.tables[0].columns[1].width)
  d2=word.export(source,self.root/'drop.docx',drop_first_title=True);self.assertEqual(d2['docxContentCounts']['Header'],1)
 def test_word_missing_dependency(self):
  source=self.write('x.md','# 标题')
  with self.assertRaisesRegex(ValueError,'Pandoc'):word.export(source,self.root/'x.docx',pandoc='nonexistent-prototype-pandoc')
if __name__=='__main__':unittest.main()
