#!/usr/bin/env python3
"""Export Markdown PRD through Pandoc, then style and structurally validate DOCX.
Requires pandoc and python-docx. Does not silently install or download anything.
"""
from __future__ import annotations
import argparse,json,math,os,re,shutil,subprocess,tempfile,zipfile
from pathlib import Path
PALETTES={'blue':('123A7A','E8F0FB','B9C9E3'),'classic-blue':('1E4F9F','EDF4FF','C3D3EC'),'light-blue':('2F6FCA','F3F8FF','D2DEF0')}

def export(source,output,pandoc='pandoc',reference_doc=None,resources=(),font='Microsoft YaHei',latin_font='Aptos',palette='blue',drop_first_title=False):
    try:
        from docx import Document
        from docx.shared import Pt,Twips,RGBColor
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
    except ImportError as e:raise ValueError('Missing python-docx; select a Python environment containing it.') from e
    binary=shutil.which(pandoc)
    if not binary:raise ValueError('Missing Pandoc; provide its executable with --pandoc.')
    source=Path(source).resolve();output=Path(output).resolve()
    if not source.is_file():raise ValueError('Markdown source not found')
    if output.exists():raise ValueError('Output exists; choose a fresh output file')
    if output.suffix.lower()!='.docx':raise ValueError('Output must be .docx')
    if reference_doc and not Path(reference_doc).is_file():raise ValueError('Reference template not found')
    color,fill,border=PALETTES[palette]
    def el(name,**attrs):
        node=OxmlElement('w:'+name)
        for k,v in attrs.items():node.set(qn('w:'+k),str(v))
        return node
    def set_font(style,size=None):
        style.font.name=latin_font
        if size:style.font.size=Pt(size)
        rpr=style.element.get_or_add_rPr();rf=rpr.find(qn('w:rFonts'))
        if rf is None:rf=el('rFonts');rpr.insert(0,rf)
        rf.set(qn('w:eastAsia'),font)
    output.parent.mkdir(parents=True,exist_ok=True)
    with tempfile.TemporaryDirectory(prefix='prd-docx-') as temp:
        temp=Path(temp);draft=temp/'draft.docx';filter_path=temp/'clean.lua'
        filter_path.write_text('''local skipped=false
function Header(el)
  if DROP and not skipped and el.level==1 then skipped=true;return {} end
  return el
end
function Image(el)
  local c=pandoc.utils.stringify(el.caption)
  if c:match("^image%-%d+$") then el.caption={} end
  return el
end
'''.replace('DROP','true' if drop_first_title else 'false'),encoding='utf-8')
        base=[binary,str(source),'--from=markdown+pipe_tables+raw_html-implicit_figures','--lua-filter='+str(filter_path),'--resource-path='+os.pathsep.join([str(source.parent),*[str(Path(p).resolve()) for p in resources]])]
        ast_run=subprocess.run([*base,'--to=json'],capture_output=True,text=True,check=True)
        ast=json.loads(ast_run.stdout)
        def count(node,kind):
            if isinstance(node,dict):return int(node.get('t')==kind)+sum(count(x,kind) for x in node.values())
            if isinstance(node,list):return sum(count(x,kind) for x in node)
            return 0
        expected={k:count(ast,k) for k in ('Header','Table','Image')}
        args=[*base,'--to=docx','--standalone','--fail-if-warnings','-o',str(draft)]
        if reference_doc:args+=['--reference-doc='+str(Path(reference_doc).resolve())]
        subprocess.run(args,capture_output=True,text=True,check=True)
        doc=Document(draft)
        for section in doc.sections:
            section.page_width=Twips(11906);section.page_height=Twips(16838);section.top_margin=Twips(1200);section.bottom_margin=Twips(1200);section.left_margin=Twips(900);section.right_margin=Twips(900)
        normal=doc.styles['Normal'];set_font(normal,11);normal.font.color.rgb=RGBColor.from_string('333333');normal.paragraph_format.space_after=Pt(5);normal.paragraph_format.line_spacing=1.3
        for i,size in enumerate([20,16,13,12,11,11],1):
            style=doc.styles['Heading '+str(i)];set_font(style,size);style.font.color.rgb=RGBColor.from_string(color);style.font.bold=True;style.paragraph_format.keep_with_next=True
        for table in doc.tables:
            table.autofit=False;pr=table._tbl.tblPr
            for name in ('tblBorders','tblW','tblLayout','tblCellMar'):
                for old in pr.findall(qn('w:'+name)):pr.remove(old)
            pr.append(el('tblW',w=5000,type='pct'));pr.append(el('tblLayout',type='fixed'))
            borders=el('tblBorders')
            for name in ('top','left','bottom','right','insideH','insideV'):borders.append(el(name,val='single',sz=4,color=border))
            pr.append(borders);mar=el('tblCellMar')
            for name in ('top','bottom','left','right'):mar.append(el(name,w=60,type='dxa'))
            pr.append(mar)
            n=len(table.columns)
            lengths=[1.0]*n
            for row in table.rows:
                for i,cell in enumerate(row.cells):
                    if i<n:lengths[i]=max(lengths[i],sum(2 if ord(c)>127 else .5 if c.isspace() else 1 for c in cell.text))
            weights=[math.sqrt(min(v,160)) for v in lengths];minimum=min(720,10106//max(n,1));remaining=10106-minimum*n
            widths=[minimum+int(remaining*w/sum(weights)) for w in weights];widths[-1]+=10106-sum(widths)
            grid=table._tbl.tblGrid
            for item in list(grid):grid.remove(item)
            for width in widths:grid.append(el('gridCol',w=width))
            for ri,row in enumerate(table.rows):
                if ri==0:
                    trpr=row._tr.get_or_add_trPr();trpr.append(el('tblHeader'))
                for ci,cell in enumerate(row.cells):
                    cell.width=Twips(widths[min(ci,n-1)]);cp=cell._tc.get_or_add_tcPr()
                    for old in cp.findall(qn('w:noWrap'))+cp.findall(qn('w:tcBorders')):cp.remove(old)
                    if ri==0:
                        for old in cp.findall(qn('w:shd')):cp.remove(old)
                        cp.append(el('shd',val='clear',fill=fill))
                    for paragraph in cell.paragraphs:
                        paragraph.paragraph_format.space_after=Pt(1);paragraph.paragraph_format.line_spacing=1.1
                        for run in paragraph.runs:
                            set_font(run,9)
                            if ri==0:run.bold=True;run.font.color.rgb=RGBColor.from_string(color)
        for node in doc.element.iter():
            if node.tag.endswith('}docPr') or node.tag.endswith('}cNvPr'):
                for key in ('descr','name'):
                    if re.search(r'image-\d+|typora-user-images',node.get(key,'')):node.set(key,'' if key=='descr' else 'Picture')
        doc.save(draft)
        # Validate emitted OOXML rather than assuming our styling calls took effect.
        verified=Document(draft);headings=sum(1 for p in verified.paragraphs if p.style.name.startswith('Heading '))
        actual={'Header':headings,'Table':len(verified.tables),'Image':len(verified.inline_shapes)}
        if expected!=actual:raise ValueError(f'Content structure mismatch: expected {expected}, got {actual}. Complex/nested structures require document-specific export.')
        for table in verified.tables:
            b=table._tbl.tblPr.find(qn('w:tblBorders'))
            if b is None or any(b.find(qn('w:'+x)) is None for x in ('top','left','bottom','right','insideH','insideV')):raise ValueError('Table grid validation failed')
            grid=[int(x.get(qn('w:w'))) for x in table._tbl.tblGrid]
            if sum(grid)!=10106:raise ValueError('Table width total invalid')
            for i,c in enumerate(table.rows[0].cells):
                if int(c._tc.tcPr.find(qn('w:tcW')).get(qn('w:w')))!=grid[i]:raise ValueError('Cell/grid width mismatch')
        with zipfile.ZipFile(draft) as z:
            if z.testzip():raise ValueError('DOCX ZIP corruption')
        shutil.copy2(draft,output)
    report=dict(output=str(output),source=str(source),palette=palette,sourceContentCounts=expected,docxContentCounts=actual,structureVerified=True,visualVerified=False,dependencies=dict(pandoc=binary,pythonDocx=True))
    report_path=output.with_suffix('.validation.json');report_path.write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    return report

def main():
    p=argparse.ArgumentParser(description=__doc__);p.add_argument('source');p.add_argument('output');p.add_argument('--pandoc',default='pandoc');p.add_argument('--reference-doc');p.add_argument('--resource-path',action='append',default=[]);p.add_argument('--font',default='Microsoft YaHei');p.add_argument('--latin-font',default='Aptos');p.add_argument('--palette',choices=PALETTES,default='blue');p.add_argument('--drop-first-title',action='store_true');a=p.parse_args()
    try:print(json.dumps(export(a.source,a.output,a.pandoc,a.reference_doc,a.resource_path,a.font,a.latin_font,a.palette,a.drop_first_title),ensure_ascii=False,indent=2))
    except (ValueError,OSError,subprocess.CalledProcessError) as e:p.exit(2,str(e)+'\n'+(getattr(e,'stderr','') or ''))
if __name__=='__main__':main()
